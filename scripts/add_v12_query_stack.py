from copy import deepcopy
from pathlib import Path

from docx import Document
from docx.oxml import OxmlElement
from docx.text.paragraph import Paragraph


RESUME = Path(r"E:\简历\宋铁轶_简历_AI应用开发_v12(单页).docx")
SCENARIO_PREFIX = "• 场景与目标："
STACK = "• 技术栈：Python / FastAPI / LangChain / DeepSeek / ChromaDB / RediSearch / Redis / RabbitMQ / Spring Boot / MyBatis / MCP / Docker"


def main():
    document = Document(RESUME)
    if any(p.text == STACK for p in document.paragraphs):
        raise RuntimeError("Query tech stack already exists.")
    scenario = next(p for p in document.paragraphs if p.text.startswith(SCENARIO_PREFIX))
    label_reference = next(p for p in document.paragraphs if p.text.startswith("• 验证范围："))

    new_element = OxmlElement("w:p")
    scenario._p.addnext(new_element)
    new_paragraph = Paragraph(new_element, scenario._parent)
    new_element.insert(0, deepcopy(scenario._p.pPr))
    label, body = STACK.split("：", 1)
    label_run = new_paragraph.add_run(label + "：")
    label_run._r.insert(0, deepcopy(label_reference.runs[0]._r.rPr))
    body_run = new_paragraph.add_run(body)
    body_run._r.insert(0, deepcopy(label_reference.runs[-1]._r.rPr))
    document.save(RESUME)


if __name__ == "__main__":
    main()
